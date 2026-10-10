"""Ingestion d'un lot "Suivi_dechet" (images + constats) vers le dashboard UVE Tibi.

Un lot est le dossier produit par le tri manuel : un sous-dossier par plaque
contenant les images, plus le tableur de suivi. Le script ecrit aux deux
endroits que lit le dashboard :

  - les images vers `tibilevel_test_isapi/Galerie_erreur/{jour}/{plaque}/` ;
  - une ligne BigQuery par couple (plaque, type), count somme, comme les lots
    precedents.

Usage :

    python ingerer_suivi.py <dossier_lot> <tableur> [--feuille NOM]
                            [--avec-photo-plaque] [--execute]

Le tableur est un .xlsx (necessite openpyxl, avec --feuille) ou un .csv
exporte depuis celui-ci. Les colonnes sont lues par position :
Date, Heure, Plaque, Dechet, Count, Nom_img.

Sans --execute le script n'ecrit rien : il affiche ce qu'il ferait. A lancer
depuis ce dossier (il importe config.py).

Ce qui est publie est exactement ce que montre la version locale du dashboard :
  - decisions de la fenetre de decision (corrections_plaques.json) : lignes du
    tableur ignorees, plaques corrigees, images retirees d'une galerie ;
  - constats sans immatriculation (« X ») et camions sans photo de plaque
    publies (conserves en base) : c'est le dashboard qui les masque ;
  - photos de plaque (« plaque_... ») publiees sous le nom « Plaque_immat_... »
    de la galerie existante, ce qui les place en tete de galerie.
"""

from __future__ import annotations

import argparse
import csv
import re
from datetime import datetime
from pathlib import Path

import config as cfg

EXTENSIONS_IMAGE = (".png", ".jpg", ".jpeg", ".webp")
PREFIXE_PHOTO_PLAQUE = "datasets_dataset_nouveau_dataset_trie_"
HEURE_RE = re.compile(r"^(\d{1,2})h(\d{1,2})m(?:in)?(\d{1,2})s$")  # 12h00min28s ou 11h56m07s
# 09:43:54 (cellule heure d'Excel) ; aussi la fin de « 1900-01-02 13:27:26 » ou « 2 days, 13:27:26 » quand la
# cellule heure a ete etiree dans Excel (la date parasite s'incremente) : seule l'heure compte.
HEURE_HMS_RE = re.compile(r"(\d{1,2}):(\d{2}):(\d{2})(?:\.\d+)?$")
FORMATS_DATE = ("%Y_%m_%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y")  # 2026_09_03 ou cellule date d'Excel
SCORE_CONFIANCE = 100  # constats relus a la main : confiance pleine


def _cellules(chemin: Path, feuille: str | None) -> list[tuple]:
    """Lignes brutes du tableur, en-tete exclue, colonnes lues par position."""
    if chemin.suffix.lower() == ".csv":
        with open(chemin, encoding="utf-8-sig", newline="") as fichier:
            return list(csv.reader(fichier))[1:]

    try:
        import openpyxl
    except ModuleNotFoundError:
        raise SystemExit(
            "Lecture .xlsx impossible : openpyxl n'est pas installe dans cet "
            "environnement.\nSoit `pip install openpyxl`, soit exportez la "
            "feuille en .csv et passez ce .csv au script."
        ) from None

    classeur = openpyxl.load_workbook(chemin, data_only=True)
    if feuille is None:
        raise SystemExit(
            f"Precisez --feuille. Feuilles disponibles : {classeur.sheetnames}"
        )
    if feuille not in classeur.sheetnames:
        raise SystemExit(
            f"Feuille '{feuille}' absente. Feuilles disponibles : {classeur.sheetnames}"
        )
    return list(classeur[feuille].iter_rows(min_row=2, values_only=True))


def lire_constats(chemin: Path, feuille: str | None) -> list[dict]:
    """Constats du tableur, lignes de separation entre vehicules ignorees."""
    lignes = []
    for numero, cellules in enumerate(_cellules(chemin, feuille), start=2):   # numero de ligne du tableur
        valeurs = [
            "" if cellule is None else str(cellule).strip()
            for cellule in (list(cellules) + [None] * 6)[:6]
        ]
        date_brute, heure, plaque, dechet, count, nom_img = valeurs
        if not (heure and plaque and dechet):
            continue
        lignes.append({
            "date": date_brute,
            "heure": heure,
            "plaque": plaque,
            "type": cfg.normaliser_type(dechet),
            "count": int(float(count)) if count else 1,
            "nom_img": nom_img or None,
            "ligne": numero,
        })
    return lignes


def horodatage(date_brute: str, heure: str) -> datetime:
    """'2026_09_03' + '12h00min28s' -> datetime naif (convention des lots precedents).

    Les tableurs generes depuis CVAT portent de vraies cellules date et heure
    d'Excel, lues ici comme '2026-09-07 00:00:00' et '09:43:54'.
    """
    correspondance = HEURE_RE.match(heure) or HEURE_HMS_RE.search(heure)
    if not correspondance:
        raise SystemExit(f"Heure illisible : {heure!r}")
    h, m, s = (int(v) for v in correspondance.groups())
    for format_date in FORMATS_DATE:
        try:
            return datetime.strptime(date_brute, format_date).replace(hour=h, minute=m, second=s)
        except ValueError:
            continue
    raise SystemExit(f"Date illisible : {date_brute!r}")


def agreger(lignes: list[dict]) -> list[dict]:
    """Une ligne BigQuery par (plaque, type), counts sommes.

    Quand le tableur horodate chaque constat a la seconde, on retient le
    premier instant du couple : il situe le passage du vehicule.
    """
    cumul: dict[tuple[str, str], dict] = {}
    for ligne in lignes:
        cle = (ligne["plaque"], ligne["type"])
        instant = horodatage(ligne["date"], ligne["heure"])
        if cle not in cumul:
            cumul[cle] = {
                "horodatage": instant,
                "plaque": ligne["plaque"],
                "type": ligne["type"],
                "count": 0,
                "score_confiance": SCORE_CONFIANCE,
            }
        cumul[cle]["horodatage"] = min(cumul[cle]["horodatage"], instant)
        cumul[cle]["count"] += ligne["count"]
    return sorted(cumul.values(), key=lambda r: (r["plaque"], r["type"]))


def recaler_sur_dossier(lignes: list[dict], dossier_lot: Path) -> list[str]:
    """Aligne les constats sur le dossier du lot, sans modifier le tableur. Renvoie les corrections faites.

    - Date : un dossier nomme JJ-MM-AAAA fait foi (un tableur copie d'une autre journee garde l'ancienne date).
    - Plaque : quand l'image d'une ligne est rangee dans le dossier d'une autre plaque ET que la plaque du tableur
      n'a aucun dossier ce jour-la (faute de frappe sur un chiffre), la ligne prend le nom du dossier,
      sinon le constat n'aurait pas d'image. Si la plaque du tableur a son propre dossier, ce sont deux camions
      differents : rien n'est change, le conflit est seulement signale.
    """
    notes = []
    try:
        jour = datetime.strptime(dossier_lot.name, "%d-%m-%Y").strftime("%Y_%m_%d")
    except ValueError:
        jour = None
    rangement: dict[str, set[str]] = {}                       # une meme image peut etre copiee dans deux dossiers
    for f in dossier_lot.glob("*/*"):
        if f.suffix.lower() in EXTENSIONS_IMAGE and not f.parent.name.startswith("_"):
            rangement.setdefault(souche_image(f.name), set()).add(f.parent.name)
    dossiers_plaques = set().union(*rangement.values()) if rangement else set()
    dates_fausses, plaques, conflits = 0, {}, {}
    for ligne in lignes:
        if jour and ligne["date"][:10].replace("-", "_") != jour:
            ligne["date"] = jour
            dates_fausses += 1
        dossiers = rangement.get(souche_image(ligne["nom_img"] or ""), set())
        # Plaque inconnue (« X ») : le constat reste sans immatriculation, donc hors vue client, meme si son image
        # a ete rangee dans le dossier d'un camion.
        if dossiers and ligne["plaque"] not in dossiers and cfg.plaque_identifiee(ligne["plaque"]):
            dossier = sorted(dossiers)[0]
            cle = (ligne["plaque"], dossier)
            if ligne["plaque"] in dossiers_plaques:
                conflits[cle] = conflits.get(cle, 0) + 1
            else:
                plaques[cle] = plaques.get(cle, 0) + 1
                ligne["plaque"] = dossier
    if dates_fausses:
        notes.append(f"date du tableur corrigee sur {dates_fausses} ligne(s) -> {dossier_lot.name}")
    notes += [f"plaque {a!r} -> {b!r} ({n} ligne(s), d'apres le dossier de l'image)" for (a, b), n in plaques.items()]
    notes += [f"CONFLIT plaque {a!r} : {n} image(s) rangee(s) dans le dossier {b!r} (laisse tel quel)"
              for (a, b), n in conflits.items()]
    return notes


def images_par_plaque(dossier_lot: Path, avec_photo_plaque: bool) -> dict[str, list[Path]]:
    resultat: dict[str, list[Path]] = {}
    # Les dossiers techniques (_corbeille, _sauvegardes, _originaux_couleurs, _originaux_floutage...) gardent des
    # copies de travail ou des originaux non floutes : jamais publies.
    for sous_dossier in sorted(p for p in dossier_lot.iterdir() if p.is_dir() and not p.name.startswith("_")):
        fichiers = [
            fichier
            for fichier in sorted(sous_dossier.iterdir())
            if fichier.suffix.lower() in EXTENSIONS_IMAGE
            and (avec_photo_plaque or not fichier.name.startswith(PREFIXE_PHOTO_PLAQUE))
        ]
        if fichiers:
            resultat[sous_dossier.name] = fichiers
    return resultat


def souche_image(reference: str) -> str:
    """Nom d'image sans dossier ni extension.

    Le tableur cite parfois un chemin CVAT en .jpg la ou le lot livre un .png :
    seule la souche du nom permet de rapprocher les deux.
    """
    nom = Path(str(reference).strip()).name.lower()
    while nom.rsplit(".", 1)[-1] in ("jpg", "jpeg", "png", "webp") and "." in nom:    # x.jpg.png -> x
        nom = nom.rsplit(".", 1)[0]
    return nom.rstrip(".").removeprefix("surveillance_")


def nom_blob(fichier: Path) -> str:
    """Le prefixe technique du dataset est retire, comme pour le lot du 31/08. Une photo de plaque « plaque_... »
    prend le nom « Plaque_immat_... » des galeries deja publiees."""
    nom = fichier.name.removeprefix(PREFIXE_PHOTO_PLAQUE)
    if nom.lower().startswith("plaque_") and not nom.lower().startswith("plaque_immat"):
        nom = "Plaque_immat_" + nom[len("plaque_"):]
    return nom


def main() -> None:
    parseur = argparse.ArgumentParser()
    parseur.add_argument("dossier_lot", type=Path)
    parseur.add_argument("tableur", type=Path)
    parseur.add_argument("--feuille", default=None)
    parseur.add_argument("--execute", action="store_true")
    parseur.add_argument("--avec-photo-plaque", action="store_true")
    args = parseur.parse_args()

    lignes = lire_constats(args.tableur, args.feuille)
    try:
        jour_dossier = datetime.strptime(args.dossier_lot.name, "%d-%m-%Y").date().isoformat()
    except ValueError:                                           # lot non nomme JJ-MM-AAAA : pas de decision
        jour_dossier = None
    ignorees = cfg.lignes_ignorees().get(jour_dossier, set())
    if ignorees:
        print(f"  (i) lignes du tableur ignorees (fenetre de decision) : {sorted(ignorees)}")
        lignes = [ligne for ligne in lignes if ligne["ligne"] not in ignorees]
    for note in recaler_sur_dossier(lignes, args.dossier_lot):
        print(f"  (i) {note}")

    images = images_par_plaque(args.dossier_lot, args.avec_photo_plaque)
    masquees = cfg.images_masquees()
    images = {
        dossier: [f for f in fichiers if f.name not in masquees.get((jour_dossier, dossier), set())]
        for dossier, fichiers in images.items()
    }
    images = {d: f for d, f in images.items() if f}

    # Publies mais masques par le dashboard (conserves en base, comme le type Textile) : un camion qui recoit plus
    # tard la photo de sa plaque reapparait sans nouvelle ingestion.
    sans_plaque = sorted({ligne["plaque"] for ligne in lignes if not cfg.plaque_identifiee(ligne["plaque"])})
    if sans_plaque:
        print(f"  (i) sans immatriculation, publie(s) mais masque(s) par le dashboard : {sans_plaque}")
    sans_photo = sorted(d for d, fichiers in images.items() if not any(cfg.est_photo_plaque(f.name) for f in fichiers))
    if sans_photo:
        print(f"  (i) sans photo de plaque, publie(s) mais masque(s) par le dashboard : {sans_photo}")

    corrections = cfg.corrections_plaques()
    renommes = {d: corrections[(jour_dossier, d)]["plaque_corrigee"] for d in images if (jour_dossier, d) in corrections}
    for ligne in lignes:
        correction = corrections.get((jour_dossier, ligne["plaque"]))
        if correction:
            ligne["plaque"] = correction["plaque_corrigee"]
    for ancien, nouveau in renommes.items():
        print(f"  (i) plaque corrigee (fenetre de decision) : {ancien} -> {nouveau}")
        images.setdefault(nouveau, []).extend(images.pop(ancien))
    if not lignes:
        raise SystemExit("Aucun constat lisible dans le tableur.")

    constats = agreger(lignes)
    jour = constats[0]["horodatage"].date()

    # -- Coherence entre le tableur et les dossiers d'images -------------------
    plaques_excel = {c["plaque"] for c in constats}
    plaques_images = set(images)
    if plaques_excel - plaques_images:
        print(f"  /!\\ Plaques du tableur sans dossier d'images : {sorted(plaques_excel - plaques_images)}")
    if plaques_images - plaques_excel:
        print(f"  /!\\ Dossiers d'images sans constat au tableur : {sorted(plaques_images - plaques_excel)}")

    souches_disque = {souche_image(f.name) for fichiers in images.values() for f in fichiers}
    manquantes = {souche_image(l["nom_img"]) for l in lignes if l["nom_img"]} - souches_disque
    if manquantes:
        print(f"  /!\\ Images citees au tableur mais absentes du lot : {sorted(manquantes)}")

    # -- Etat actuel en production --------------------------------------------
    df = cfg.charger_constats()
    deja_bq = len(df[df["date"] == jour])
    deja_gcs = {
        b.name for b in cfg.client_storage().bucket(cfg.BUCKET).list_blobs(
            prefix=f"{cfg.GALERIE_PREFIX}/{jour.isoformat()}/"
        )
    }

    print(f"\n=== Lot du {jour.isoformat()} ===")
    print(f"Tableur   : {len(lignes)} lignes detaillees -> {len(constats)} constats agreges")
    print(f"Images    : {sum(len(v) for v in images.values())} fichiers sur {len(images)} plaques")
    print(f"Existant  : {deja_bq} ligne(s) BigQuery et {len(deja_gcs)} objet(s) GCS pour ce jour")

    print("\n--- Lignes BigQuery a inserer ---")
    for c in constats:
        print(f"  {c['horodatage']}  {c['plaque']:<10} {c['type']:<14} count={c['count']:<3} conf={c['score_confiance']}")

    print("\n--- Images a envoyer ---")
    total_octets = 0
    for plaque, fichiers in images.items():
        octets = sum(f.stat().st_size for f in fichiers)
        total_octets += octets
        print(f"  {plaque:<10} {len(fichiers):>2} fichier(s)  {octets / 1024 / 1024:6.1f} Mo")
    print(f"  {'TOTAL':<10} {sum(len(v) for v in images.values()):>2} fichier(s)  {total_octets / 1024 / 1024:6.1f} Mo")

    if not args.execute:
        print("\n[DRY-RUN] Rien n'a ete ecrit. Relancer avec --execute pour appliquer.")
        return

    if deja_bq:
        raise SystemExit(
            f"\nARRET : {deja_bq} ligne(s) existent deja en BigQuery pour le {jour}. "
            "Supprimez-les d'abord (page Administration) pour eviter les doublons."
        )

    # -- Ecriture : images d'abord, constats ensuite ---------------------------
    print("\n--- Envoi des images ---")
    bucket = cfg.client_storage().bucket(cfg.BUCKET)
    envoyees = 0
    for plaque, fichiers in images.items():
        for fichier in fichiers:
            chemin = f"{cfg.GALERIE_PREFIX}/{jour.isoformat()}/{plaque}/{nom_blob(fichier)}"
            if chemin in deja_gcs:
                print(f"  = deja present : {chemin}")
                continue
            bucket.blob(chemin).upload_from_filename(str(fichier))
            envoyees += 1
            print(f"  + {chemin}")
    print(f"{envoyees} image(s) envoyee(s).")

    print("\n--- Insertion des constats ---")
    erreurs = cfg.client_bigquery().insert_rows_json(
        cfg.TABLE_REF,
        [{
            "horodatage": c["horodatage"].isoformat(),
            "type": c["type"],
            "count": c["count"],
            "score_confiance": c["score_confiance"],
            "plaque": c["plaque"],
        } for c in constats],
    )
    if erreurs:
        raise SystemExit(f"Echec d'insertion BigQuery : {erreurs}")
    print(f"{len(constats)} constat(s) inseres dans {cfg.TABLE_REF}.")
    print("\nTermine. Le dashboard affichera le lot d'ici 5 minutes (cache) ou apres redemarrage.")


if __name__ == "__main__":
    main()
